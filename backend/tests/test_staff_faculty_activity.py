"""Faculty Activity Report — teacher biometric attendance joined with the
lectures they delivered (scheduled vs actual, delays)."""

import io
import uuid
from datetime import date, datetime, timezone

from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.attendance.models.attendance_models import StaffDailyAttendance
from app.modules.attendance.services import staff_faculty_service as fac
from app.modules.lectures.models.lecture_models import Lecture
from app.modules.staff.models.staff_models import Department, Staff

BRANCH_A = uuid.UUID("00000000-0000-0000-0000-000000000001")
TEACHER = uuid.UUID("00000000-0000-0000-0000-000000000060")
BATCH = uuid.UUID("00000000-0000-0000-0000-000000000070")
SUBJECT = uuid.UUID("00000000-0000-0000-0000-000000000050")
AY = uuid.UUID("00000000-0000-0000-0000-000000000030")
DAY = date(2026, 8, 3)


def _utc(h, mi):
    return datetime(2026, 8, 3, h, mi, tzinfo=timezone.utc)


async def _seed(db: AsyncSession) -> Staff:
    dept = Department(branch_id=BRANCH_A, name="MSA-Teachers", id_range_start=1, id_range_end=50)
    db.add(dept)
    await db.flush()
    staff = Staff(
        branch_id=BRANCH_A, emp_code="6", first_name="Ajinkya", last_name="Deshpande",
        department_id=dept.id, linked_teacher_id=TEACHER,
    )
    db.add(staff)
    await db.flush()

    db.add(StaffDailyAttendance(
        staff_id=staff.id, branch_id=BRANCH_A, attendance_date=DAY,
        first_in=_utc(5, 30), last_out=_utc(12, 30), day_status="PRESENT",
        signoff="COMPLETE", source="BIOMETRIC", work_minutes=420, ot_minutes=0,
    ))

    # Two lectures: one on time, one delayed 15 min (late_flag set).
    db.add(Lecture(
        teacher_id=TEACHER, batch_id=BATCH, subject_id=SUBJECT,
        scheduled_start=_utc(8, 30), scheduled_end=_utc(9, 30),
        actual_start=_utc(8, 30), actual_end=_utc(9, 30),
        actual_duration_min=60, late_flag=False, lecture_status="completed",
        branch_id=BRANCH_A, academic_year_id=AY,
    ))
    db.add(Lecture(
        teacher_id=TEACHER, batch_id=BATCH, subject_id=SUBJECT,
        scheduled_start=_utc(9, 30), scheduled_end=_utc(10, 30),
        actual_start=_utc(9, 45), actual_end=_utc(10, 30),
        actual_duration_min=45, late_flag=True, lecture_status="completed",
        branch_id=BRANCH_A, academic_year_id=AY,
    ))
    await db.flush()
    return staff


def _cells(data: bytes) -> dict:
    ws = load_workbook(io.BytesIO(data), read_only=True).active
    return {r[0]: r[1] for r in ws.iter_rows(values_only=True) if r[0]}


async def test_daily_faculty_activity(db_session: AsyncSession, seed_data):
    await _seed(db_session)
    fn, data, mime = await fac.daily_report(
        db_session, branch_id=BRANCH_A, teacher_id=TEACHER, day=DAY, fmt="xlsx",
    )
    assert mime == fac.XLSX_MIME
    c = _cells(data)
    assert c["Total Lectures"] == 2
    assert c["Delayed Lectures"] == 1
    assert c["Time Lost to Delay"] == "0:15"      # 15-min delay on the 2nd lecture
    assert c["Scheduled Lecture Hours"] == "2:00"  # two 1h lectures
    assert c["Effective Lecture Hours"] == "1:45"  # 60 + 45 min
    assert c["Attendance Hours"] == "7:00"         # 420 min work


async def test_daily_faculty_activity_lists_delayed_lectures(
    db_session: AsyncSession, seed_data
):
    """The Excel export carries the Delayed Lectures detail table (Section 2)."""
    await _seed(db_session)
    _, data, _ = await fac.daily_report(
        db_session, branch_id=BRANCH_A, teacher_id=TEACHER, day=DAY, fmt="xlsx",
    )
    ws = load_workbook(io.BytesIO(data), read_only=True).active
    col_a = [r[0] for r in ws.iter_rows(values_only=True)]
    assert "Delayed Lecture Details" in col_a
    # The one delayed lecture (15-min delay) shows as a numbered detail row.
    rows = list(ws.iter_rows(values_only=True))
    hdr = next(i for i, r in enumerate(rows) if r[0] == "Delayed Lecture Details")
    detail = rows[hdr + 2]  # title, header, then first data row
    assert detail[0] == 1
    assert detail[5] == "0:15"  # delay


async def test_cumulative_summary(db_session: AsyncSession, seed_data):
    await _seed(db_session)
    fn, data, _ = await fac.summary_report(
        db_session, branch_id=BRANCH_A, start=DAY, end=DAY,
        teacher_ids=None, fmt="xlsx",
    )
    ws = load_workbook(io.BytesIO(data), read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    # New schema: Sr No, Employee Code, Initials, Teacher, Subject,
    # Present Days, Total Lectures, Total Hours.
    data_row = next(r for r in rows if r[1] == "6")  # emp_code in column 2
    assert data_row[0] == 1                          # Sr No
    assert data_row[2] == fac._initials(data_row[3])  # initials auto from name
    assert data_row[2] and data_row[2].isupper()      # non-empty, uppercase
    assert data_row[5] == 1                          # present days
    assert data_row[6] == 2                          # total lectures
    assert data_row[7] == "1 Hours 45 Mins"          # total (effective) hours: 60+45


def test_initials_helper():
    assert fac._initials("Bhagvat Dhesale") == "BD"
    assert fac._initials("Mr. Anish A") == "MAA"
    assert fac._initials("Madonna") == "M"
    assert fac._initials("") == ""
