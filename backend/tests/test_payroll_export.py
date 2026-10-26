"""Payroll — Slice 5 PDF/Excel exports (§7/§11).

Exercises the Excel exports end-to-end (deterministic, no browser) and the HTML
builders the PDF path renders. Comparison check: the exported numbers equal the
on-screen report (both fed by report_service), and the late fine equals the
existing shared slab logic.
"""

import io
import uuid
from datetime import date

from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.attendance.models.attendance_models import (
    LateFineSlab,
    StaffDailyAttendance,
)
from app.modules.attendance.services.daily_service import branch_timezone
from app.modules.attendance.time_utils import local_time_on
from app.modules.payroll.services import export_service, report_service
from app.modules.staff.models.staff_models import Department, Staff

BRANCH = "00000000-0000-0000-0000-000000000001"
BR = uuid.UUID(BRANCH)
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def _staff(db: AsyncSession, *, full: int, year: int, month: int, late: int = 0) -> Staff:
    dept = Department(branch_id=BR, name="MSA-Teachers", id_range_start=91000, id_range_end=91999)
    db.add(dept)
    await db.flush()
    s = Staff(branch_id=BR, emp_code="91000", first_name="Asha", last_name="Rao",
              department_id=dept.id, shift_start="08:00")
    db.add(s)
    await db.flush()
    tz = await branch_timezone(db, BR)
    for d in range(1, full + 1):
        in_dt = local_time_on(date(year, month, d), tz, 8, 20) if d <= late else None
        db.add(StaffDailyAttendance(
            branch_id=BR, staff_id=s.id, attendance_date=date(year, month, d),
            day_status="PRESENT", signoff="COMPLETE", source="BIOMETRIC", first_in=in_dt,
        ))
    await db.commit()
    return s


async def _slabs(db: AsyncSession) -> None:
    db.add_all([
        LateFineSlab(min_minutes=1, max_minutes=5, fine_amount=200, color_hex="#BDD7EE", label="1-5"),
        LateFineSlab(min_minutes=6, max_minutes=10, fine_amount=500, color_hex="#FFE699", label="6-10"),
        LateFineSlab(min_minutes=11, max_minutes=None, fine_amount=1000, color_hex="#F4B084", label="11+"),
    ])
    await db.commit()


async def _token(client, email, password):
    r = await client.post("/api/v1/auth/login", json={"email": email, "password": password})
    return r.cookies["access_token"]


async def _setup_month(client, tok, year, month, required=26):
    await client.put(
        f"/api/v1/payroll/months/{year}/{month}?branch_id={BRANCH}",
        json={"required_days": required}, cookies={"access_token": tok},
    )


# ── report export ────────────────────────────────────────────────────────────

async def test_export_report_xlsx(client, db_session, seed_data):
    await _staff(db_session, full=27, year=2026, month=6)
    tok = await _token(client, "admin@test.com", "Admin123!")
    await _setup_month(client, tok, 2026, 6)
    r = await client.get(
        f"/api/v1/payroll/report/export?branch_id={BRANCH}&year=2026&month=6&fmt=xlsx",
        cookies={"access_token": tok},
    )
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(XLSX_MIME)
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    # Header row 5, first data row 6 holds the staff code.
    assert ws.cell(row=6, column=1).value == "91000"
    # Net column (7th) is +1.
    assert ws.cell(row=6, column=7).value == "+1"


async def test_export_requires_role(client, db_session, seed_data):
    await _staff(db_session, full=26, year=2026, month=6)
    tok = await _token(client, "teacher@test.com", "Teacher123!")
    r = await client.get(
        f"/api/v1/payroll/report/export?branch_id={BRANCH}&year=2026&month=6&fmt=xlsx",
        cookies={"access_token": tok},
    )
    assert r.status_code == 403


async def test_export_logs_audit(client, db_session, seed_data):
    from sqlalchemy import select

    from app.modules.audit.models.audit_models import AuditLog

    await _staff(db_session, full=26, year=2026, month=6)
    tok = await _token(client, "admin@test.com", "Admin123!")
    await _setup_month(client, tok, 2026, 6)
    await client.get(
        f"/api/v1/payroll/report/export?branch_id={BRANCH}&year=2026&month=6&fmt=xlsx",
        cookies={"access_token": tok},
    )
    rows = (await db_session.execute(
        select(AuditLog).where(AuditLog.action == "Export Payroll Report")
    )).scalars().all()
    assert len(rows) >= 1


# ── ledger export ────────────────────────────────────────────────────────────

async def test_export_ledger_xlsx(client, db_session, seed_data):
    s = await _staff(db_session, full=27, year=2026, month=6)
    tok = await _token(client, "admin@test.com", "Admin123!")
    await _setup_month(client, tok, 2026, 6)
    await client.post(
        f"/api/v1/payroll/months/2026/6/close?branch_id={BRANCH}",
        json={}, cookies={"access_token": tok},
    )
    r = await client.get(
        f"/api/v1/payroll/staff/{s.id}/ledger/export?branch_id={BRANCH}&fmt=xlsx",
        cookies={"access_token": tok},
    )
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    assert ws.cell(row=6, column=2).value == "extra_added"
    assert ws.cell(row=6, column=3).value == "+1"


# ── comparison: export == screen, and late fine matches shared logic ─────────

async def test_export_numbers_match_report_and_late_fine(client, db_session, seed_data):
    # 26 present days, first 2 late by ~20 min (top band → Rs 1000 each).
    await _slabs(db_session)
    await _staff(db_session, full=26, year=2026, month=6, late=2)
    tok = await _token(client, "admin@test.com", "Admin123!")
    await _setup_month(client, tok, 2026, 6)

    screen = await report_service.monthly_report(
        db_session, branch_id=BR, year=2026, month=6
    )
    row = screen["rows"][0]
    assert row["late_fine_total"] == 2000.0  # two late days, 11+ band

    # The Excel export carries the same late-fine number (column 12).
    _, data, _ = await export_service.report(db_session, branch_id=BR, year=2026, month=6, fmt="xlsx")
    ws = load_workbook(io.BytesIO(data)).active
    assert ws.cell(row=6, column=12).value == "2000"


def test_report_html_builder_contains_rows():
    data = {
        "year": 2026, "month": 6, "lifecycle": "OPEN",
        "rows": [{
            "staff_id": "x", "emp_code": "91000", "name": "Asha Rao",
            "department": "MSA-Teachers", "required_days": 26, "present_days": 27,
            "half_days": 0, "paid_leave_days": 0, "net_days": 1, "balance_in": 0,
            "balance_out": 1, "uncovered_days": 0, "deferred_days": 0,
            "late_fine_total": 0,
        }],
        "totals": {"staff_count": 1, "present_days": 27, "uncovered_days": 0, "late_fine_total": 0},
    }
    html = export_service._report_html("Academy", "June 2026", "OPEN", data, "gen")
    assert "Monthly Payroll Attendance Report" in html
    assert "91000" in html and "Asha Rao" in html
    assert "+1" in html  # net column
