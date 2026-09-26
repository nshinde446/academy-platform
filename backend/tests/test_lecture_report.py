"""Lecture Report — flat per-lecture export (Section 1 of the Teacher
Productivity Report requirement): added columns, time-only schedule/actual
columns, and Duration as minutes or HH:MM."""

import io
import uuid
from datetime import date, datetime, timezone

from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.lectures.models.lecture_models import Lecture
from app.modules.lectures.services import lecture_report_service as lr

BRANCH_A = uuid.UUID("00000000-0000-0000-0000-000000000001")
TEACHER = uuid.UUID("00000000-0000-0000-0000-000000000060")
BATCH = uuid.UUID("00000000-0000-0000-0000-000000000070")
SUBJECT = uuid.UUID("00000000-0000-0000-0000-000000000050")
CLASSROOM = uuid.UUID("00000000-0000-0000-0000-000000000080")
ADMIN = uuid.UUID("00000000-0000-0000-0000-000000000100")
AY = uuid.UUID("00000000-0000-0000-0000-000000000030")
DAY = date(2026, 8, 3)


def _utc(h, mi):
    return datetime(2026, 8, 3, h, mi, tzinfo=timezone.utc)


async def _seed_lecture(db: AsyncSession):
    lec = Lecture(
        teacher_id=TEACHER, batch_id=BATCH, subject_id=SUBJECT, classroom_id=CLASSROOM,
        scheduled_start=_utc(4, 0), scheduled_end=_utc(5, 0),
        actual_start=_utc(4, 15), actual_end=_utc(5, 0),
        actual_duration_min=45, late_flag=True, lecture_status="completed",
        delivery_mode="offline", updated_by=ADMIN,
        branch_id=BRANCH_A, academic_year_id=AY,
    )
    db.add(lec)
    await db.flush()


def _sheet(data: bytes):
    ws = load_workbook(io.BytesIO(data), read_only=True).active
    return list(ws.iter_rows(values_only=True))


async def test_lecture_report_columns_and_minutes(db_session: AsyncSession, seed_data):
    await _seed_lecture(db_session)
    fn, data, mime = await lr.generate(
        db_session, branch_id=BRANCH_A, start=DAY, end=DAY, fmt="xlsx",
        duration_style="min",
    )
    assert mime == lr.XLSX_MIME
    rows = _sheet(data)
    header = next(r for r in rows if r and r[0] == "Date")
    assert list(header) == lr.HEADERS
    data_row = rows[rows.index(header) + 1]
    cells = dict(zip(lr.HEADERS, data_row))
    assert cells["Date"] == "03-Aug-2026"
    assert cells["Branch"]                      # branch name resolved, non-empty
    assert cells["Teacher"]                     # teacher name resolved
    assert cells["Duration"] == "45"            # raw minutes
    assert cells["Status"] == "Completed"
    assert cells["Attendance Updated By"]       # admin user resolved
    assert cells["Timestamp"]                   # updated_at rendered
    assert cells["Late Start"] == "15 min"      # 15-min late start


async def test_lecture_report_duration_hhmm(db_session: AsyncSession, seed_data):
    await _seed_lecture(db_session)
    _, data, _ = await lr.generate(
        db_session, branch_id=BRANCH_A, start=DAY, end=DAY, fmt="xlsx",
        duration_style="hhmm",
    )
    rows = _sheet(data)
    header = next(r for r in rows if r and r[0] == "Date")
    cells = dict(zip(lr.HEADERS, rows[rows.index(header) + 1]))
    assert cells["Duration"] == "0:45"          # HH:MM


async def test_lecture_report_teacher_filter_excludes_others(
    db_session: AsyncSession, seed_data
):
    await _seed_lecture(db_session)
    other = uuid.UUID("00000000-0000-0000-0000-0000000000ff")
    _, data, _ = await lr.generate(
        db_session, branch_id=BRANCH_A, start=DAY, end=DAY, fmt="xlsx",
        teacher_id=other,
    )
    rows = _sheet(data)
    header = next(r for r in rows if r and r[0] == "Date")
    # No data row after the header when filtered to a teacher with no lectures.
    assert rows.index(header) == len(rows) - 1
