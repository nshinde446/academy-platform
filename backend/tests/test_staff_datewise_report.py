"""Staff Monthly "Date Wise Attendance" report — one row per staff with a coded
cell per calendar day (P/A/L/HD/WO/HL) + the six totals, legend and note. Excel
path only (no headless Chromium in tests)."""

import io
import uuid
from datetime import date, datetime, timezone

import pytest
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.attendance.models.attendance_models import StaffDailyAttendance
from app.modules.attendance.services import staff_report_service as rpt
from app.modules.attendance.services import staff_export_service as ex
from app.modules.staff.models.staff_models import Department, Staff

BRANCH_A = uuid.UUID("00000000-0000-0000-0000-000000000001")


def _utc(d: date, h: int, mi: int) -> datetime:
    return datetime(d.year, d.month, d.day, h, mi, tzinfo=timezone.utc)


async def _seed(db: AsyncSession) -> Staff:
    dept = Department(branch_id=BRANCH_A, name="MSA-Teachers", id_range_start=1, id_range_end=50)
    db.add(dept)
    await db.flush()
    # weekly_off_days default (Sun=6); 2026-10 has Sundays on 4, 11, 18, 25.
    staff = Staff(
        branch_id=BRANCH_A, emp_code="E001", first_name="Staff", last_name="A",
        department_id=dept.id, shift_start="10:00", shift_end="20:00",
        weekly_off_days="6",
    )
    db.add(staff)
    await db.flush()

    # Oct 1 PRESENT, Oct 2 LATE, Oct 3 HALF_DAY, Oct 5 ABSENT (explicit), Oct 6
    # HOLIDAY. Oct 4 has no row and is a Sunday -> WO. Other days no row -> A.
    rows = [
        (date(2026, 10, 1), "PRESENT"),
        (date(2026, 10, 2), "LATE"),
        (date(2026, 10, 3), "HALF_DAY"),
        (date(2026, 10, 5), "ABSENT"),
        (date(2026, 10, 6), "HOLIDAY"),
    ]
    for d, st in rows:
        db.add(StaffDailyAttendance(
            staff_id=staff.id, branch_id=BRANCH_A, attendance_date=d,
            first_in=_utc(d, 4, 30) if st not in ("ABSENT", "HOLIDAY") else None,
            last_out=None, day_status=st, signoff="COMPLETE",
            source="BIOMETRIC", work_minutes=0, ot_minutes=0,
        ))
    await db.flush()
    return staff


def _ws(data: bytes):
    return load_workbook(io.BytesIO(data), read_only=True).active


async def test_datewise_full_month_header_and_columns(db_session: AsyncSession, seed_data):
    await _seed(db_session)
    fn, data, mime = await rpt.generate_datewise(
        db_session, branch_id=BRANCH_A, department_ids=None, staff_ids=None,
        start=date(2026, 10, 1), end=date(2026, 10, 31), fmt="xlsx",
    )
    assert fn.endswith(".xlsx")
    assert mime == rpt.XLSX_MIME
    ws = _ws(data)
    rows = list(ws.iter_rows(values_only=True))
    col_a = [r[0] for r in rows]
    # Full calendar month -> "Month - October 2026" banner.
    assert "Month - October 2026" in col_a
    # Header row: Empcode, Name, 31 day columns, then 6 totals = 39 columns.
    header = next(r for r in rows if r[0] == "Empcode")
    assert header[1] == "Name"
    assert header[2] == "1" and header[32] == "31"  # day columns 1..31
    assert list(header[33:39]) == ex.DATEWISE_TOTAL_CODES  # P A HD L HL WO


async def test_datewise_codes_and_totals_reconcile(db_session: AsyncSession, seed_data):
    await _seed(db_session)
    _, data, _ = await rpt.generate_datewise(
        db_session, branch_id=BRANCH_A, department_ids=None, staff_ids=None,
        start=date(2026, 10, 1), end=date(2026, 10, 31), fmt="xlsx",
    )
    ws = _ws(data)
    rows = list(ws.iter_rows(values_only=True))
    data_row = next(r for r in rows if r[0] == "E001")
    days = list(data_row[2:33])  # 31 day cells
    assert days[0] == "P"      # Oct 1 present
    assert days[1] == "L"      # Oct 2 late
    assert days[2] == "HD"     # Oct 3 half day
    assert days[3] == "WO"     # Oct 4 Sunday, no row
    assert days[4] == "A"      # Oct 5 absent
    assert days[5] == "HL"     # Oct 6 holiday

    totals = dict(zip(ex.DATEWISE_TOTAL_CODES, data_row[33:39]))
    # Totals reconcile with the day cells.
    for code in ex.DATEWISE_TOTAL_CODES:
        assert totals[code] == days.count(code), code
    assert totals["P"] == 1 and totals["L"] == 1 and totals["HD"] == 1
    assert totals["HL"] == 1
    # Sundays 4/11/18/25 are WO; everything else with no row is A.
    assert totals["WO"] == 4
    # Every day carries exactly one code: the six totals sum to 31.
    assert sum(totals.values()) == 31


async def test_datewise_late_cell_is_red_filled(db_session: AsyncSession, seed_data):
    await _seed(db_session)
    _, data, _ = await rpt.generate_datewise(
        db_session, branch_id=BRANCH_A, department_ids=None, staff_ids=None,
        start=date(2026, 10, 1), end=date(2026, 10, 31), fmt="xlsx",
    )
    wb = load_workbook(io.BytesIO(data))  # not read_only, to read fills
    ws = wb.active
    # Find the E001 row, then the Oct-2 (L) cell at column 4 (1,2 = emp/name; day2 = col 4).
    for row in ws.iter_rows():
        if row[0].value == "E001":
            late_cell = row[3]  # Oct 2
            assert late_cell.value == "L"
            assert late_cell.fill.fgColor.rgb.endswith("F8C9C9")  # red shade
            hd_cell = row[4]  # Oct 3
            assert hd_cell.fill.fgColor.rgb.endswith("C9DAF8")    # blue shade
            break
    else:
        pytest.fail("E001 row not found")


async def test_datewise_custom_range_period_header(db_session: AsyncSession, seed_data):
    await _seed(db_session)
    _, data, _ = await rpt.generate_datewise(
        db_session, branch_id=BRANCH_A, department_ids=None, staff_ids=None,
        start=date(2026, 10, 5), end=date(2026, 10, 20), fmt="xlsx",
    )
    col_a = [r[0] for r in _ws(data).iter_rows(values_only=True)]
    assert "Period - 05 Oct 2026 to 20 Oct 2026" in col_a


async def test_datewise_rejects_overlong_range(db_session: AsyncSession, seed_data):
    await _seed(db_session)
    with pytest.raises(Exception) as exc:
        await rpt.generate_datewise(
            db_session, branch_id=BRANCH_A, department_ids=None, staff_ids=None,
            start=date(2026, 1, 1), end=date(2026, 12, 31), fmt="xlsx",
        )
    assert "maximum" in str(exc.value).lower()


async def test_datewise_rejects_from_after_to(db_session: AsyncSession, seed_data):
    await _seed(db_session)
    with pytest.raises(Exception):
        await rpt.generate_datewise(
            db_session, branch_id=BRANCH_A, department_ids=None, staff_ids=None,
            start=date(2026, 10, 20), end=date(2026, 10, 5), fmt="xlsx",
        )


async def test_datewise_legend_and_note_printed(db_session: AsyncSession, seed_data):
    await _seed(db_session)
    _, data, _ = await rpt.generate_datewise(
        db_session, branch_id=BRANCH_A, department_ids=None, staff_ids=None,
        start=date(2026, 10, 1), end=date(2026, 10, 31), fmt="xlsx",
    )
    col_a = [str(r[0]) for r in _ws(data).iter_rows(values_only=True) if r[0] is not None]
    assert "Legend" in col_a
    assert "Notes" in col_a
    assert "This is for your reference." in " ".join(col_a)
