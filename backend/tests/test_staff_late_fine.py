"""Staff Late Fine Report — fine calc + per-day rows + colored export + total."""

import io
import uuid
from datetime import date, datetime, timezone

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.attendance.models.attendance_models import (
    LateFineSlab,
    StaffDailyAttendance,
)
from app.modules.attendance.services import staff_late_fine_service as lf
from app.modules.staff.models.staff_models import Department, Staff

BRANCH = uuid.UUID("00000000-0000-0000-0000-000000000001")
UTC = timezone.utc


def _slab(mn, mx, amt, hx, lbl):
    return LateFineSlab(min_minutes=mn, max_minutes=mx, fine_amount=amt, color_hex=hx, label=lbl)


async def _seed(db: AsyncSession) -> Staff:
    db.add_all([
        _slab(1, 5, 200, "#BDD7EE", "1 to 5"),
        _slab(6, 10, 500, "#FFEB84", "6 to 10"),
        _slab(11, None, 1000, "#FF6B6B", "11 or more"),
    ])
    dept = Department(branch_id=BRANCH, name="MSA-Teachers", id_range_start=1, id_range_end=50)
    db.add(dept)
    await db.flush()
    staff = Staff(
        branch_id=BRANCH, emp_code="12", first_name="Rahul", last_name="Pawar",
        department_id=dept.id, shift_start="08:00",
    )
    db.add(staff)
    await db.flush()
    # IST = UTC+5:30, so local 08:04 = UTC 02:34 (matches the spec's example).
    for day, h, m in [(1, 2, 28), (2, 2, 34), (3, 2, 38), (4, 2, 46)]:  # 07:58, 08:04, 08:08, 08:16 IST
        db.add(StaffDailyAttendance(
            staff_id=staff.id, branch_id=BRANCH, attendance_date=date(2026, 10, day),
            first_in=datetime(2026, 10, day, h, m, tzinfo=UTC),
            day_status="PRESENT", signoff="COMPLETE", source="BIOMETRIC", work_minutes=400,
        ))
    await db.flush()
    return staff


def test_parse_hhmm_tolerant_of_separator():
    # Staff enter the First Lecture Start Time with various separators; a dot was
    # the field bug (08.00) that made every Late Minutes read 0.
    for raw in ("08:00", "08.00", "8-00", "0800", "800"):
        assert lf._parse_hhmm(raw) == (8, 0), raw
    assert lf._parse_hhmm("") is None
    assert lf._parse_hhmm(None) is None
    # A dot-separated start now measures lateness correctly.
    assert lf.late_minutes("08.00", datetime(2026, 10, 5, 8, 4)) == 4
    assert lf.late_minutes("08.00", datetime(2026, 10, 1, 18, 22)) == 622
    assert lf.late_minutes("08.00", datetime(2026, 10, 14, 7, 59)) == 0


async def test_late_fine_with_dot_separated_shift_start(
    db_session: AsyncSession, seed_data
):
    """Reproduces the field report (emp 91008): First Lecture Start Time stored
    as "08.00" (dot). Late Minutes/Fine must compute, not read 0 everywhere."""
    db_session.add_all([
        _slab(1, 5, 200, "#BDD7EE", "1 to 5"),
        _slab(6, 10, 500, "#FFEB84", "6 to 10"),
        _slab(11, None, 1000, "#FF6B6B", "11 or more"),
    ])
    dept = Department(branch_id=BRANCH, name="MSA-Teachers", id_range_start=1, id_range_end=50)
    db_session.add(dept)
    await db_session.flush()
    staff = Staff(
        branch_id=BRANCH, emp_code="91008", first_name="Deepak", last_name="Singh",
        department_id=dept.id, shift_start="08.00",  # entered with a dot
    )
    db_session.add(staff)
    await db_session.flush()
    # IST = UTC+5:30. local 08:04 = 02:34 UTC, 08:06 = 02:36, 18:22 = 12:52, 07:59 = 02:29.
    for day, h, m in [(5, 2, 34), (6, 2, 36), (1, 12, 52), (14, 2, 29)]:
        db_session.add(StaffDailyAttendance(
            staff_id=staff.id, branch_id=BRANCH, attendance_date=date(2026, 10, day),
            first_in=datetime(2026, 10, day, h, m, tzinfo=UTC),
            day_status="PRESENT", signoff="COMPLETE", source="BIOMETRIC", work_minutes=400,
        ))
    await db_session.flush()

    _, data, _ = await lf.generate(
        db_session, branch_id=BRANCH, staff_id=staff.id,
        start=date(2026, 10, 1), end=date(2026, 10, 14), fmt="xlsx",
    )
    ws = load_workbook(io.BytesIO(data)).active
    rows = list(ws.iter_rows(values_only=True))
    hdr = next(i for i, r in enumerate(rows) if r[0] == "Date")
    by_date = {r[0]: r for r in rows[hdr + 1:] if isinstance(r[0], str) and r[0].count("-") == 2}
    # start column shown as entered; late minutes + fines now computed.
    assert by_date["01-10-2026"][4] == 622 and by_date["01-10-2026"][5] == 1000  # 18:22
    assert by_date["05-10-2026"][4] == 4 and by_date["05-10-2026"][5] == 200     # 08:04
    assert by_date["06-10-2026"][4] == 6 and by_date["06-10-2026"][5] == 500     # 08:06
    assert by_date["14-10-2026"][4] == 0 and by_date["14-10-2026"][5] == 0       # 07:59 on time


def test_fine_calc_matches_spec_example():
    slabs = [_slab(1, 5, 200, "#BDD7EE", "a"), _slab(6, 10, 500, "#FFEB84", "b"),
             _slab(11, None, 1000, "#FF6B6B", "c")]
    for mins, exp_fine in [(0, 0), (4, 200), (5, 200), (8, 500), (10, 500), (11, 1000), (16, 1000)]:
        assert lf.fine_for(mins, slabs)[0] == exp_fine
    # early/on-time never negative
    assert lf.late_minutes("08:00", datetime(2026, 10, 2, 7, 58)) == 0
    assert lf.late_minutes("08:00", datetime(2026, 10, 2, 8, 16)) == 16


async def test_late_fine_xlsx_rows_colors_and_total(db_session: AsyncSession, seed_data):
    staff = await _seed(db_session)
    fn, data, mime = await lf.generate(
        db_session, branch_id=BRANCH, staff_id=staff.id,
        start=date(2026, 10, 1), end=date(2026, 10, 4), fmt="xlsx",
    )
    assert fn.startswith("LateFine_12_2026-10-01_to_2026-10-04") and mime == lf.XLSX_MIME
    ws = load_workbook(io.BytesIO(data)).active
    rows = list(ws.iter_rows(values_only=True))
    # locate the header, then the 4 data rows (Date, Start, IN, LateIN, LateMin, Fine)
    hdr = next(i for i, r in enumerate(rows) if r[0] == "Date")
    data_rows = rows[hdr + 1: hdr + 5]
    late_min = [r[4] for r in data_rows]
    fines = [r[5] for r in data_rows]
    assert late_min == [0, 4, 8, 16]
    assert fines == [0, 200, 500, 1000]
    # Late IN formatted hh:mm
    assert [r[3] for r in data_rows] == ["00:00", "00:04", "00:08", "00:16"]
    # Total Fine row
    total_row = rows[hdr + 5]
    assert total_row[3] == "Total Fine" and total_row[5] == 1700
    # Late Minutes cell is colored per slab (16 -> red #FF6B6B).
    ws2 = load_workbook(io.BytesIO(data)).active
    red_cell = ws2.cell(row=hdr + 1 + 4, column=5)  # 4th data row, col E
    assert red_cell.fill.fgColor.rgb.endswith("FF6B6B")


def test_late_fine_html_has_colored_cell_and_total():
    # The PDF is rendered from this HTML (Chromium isn't available in tests).
    report = {
        "brand": "Matrix", "emp_code": "12", "name": "Rahul Pawar",
        "period": "01-10-2026 to 04-10-2026", "total": 1700.0,
        "rows": [
            {"date": "02-10-2026", "start": "08:00", "in": "08:04", "late_in": "00:04",
             "late_minutes": 4, "fine": 200, "color": "#BDD7EE"},
            {"date": "04-10-2026", "start": "08:00", "in": "08:16", "late_in": "00:16",
             "late_minutes": 16, "fine": 1000, "color": "#FF6B6B"},
        ],
    }
    html = lf._html(report, "gen")
    assert "Staff Late Fine Report" in html
    assert "background:#FF6B6B" in html          # Late Minutes cell colored
    assert "Total Fine" in html and ">1700<" in html


def test_start_for_day_weekday_vs_saturday():
    """Institute shift: 08:00 on weekdays (Mon–Fri), 10:00 on the weekend
    (Sat/Sun), when the staff has no per-person shift set."""
    from types import SimpleNamespace
    no_shift = SimpleNamespace(shift_start=None)
    assert lf.start_for_day(date(2026, 10, 5), no_shift) == "08:00"   # Monday
    assert lf.start_for_day(date(2026, 10, 9), no_shift) == "08:00"   # Friday
    assert lf.start_for_day(date(2026, 10, 10), no_shift) == "10:00"  # Saturday
    assert lf.start_for_day(date(2026, 10, 11), no_shift) == "10:00"  # Sunday
    # A per-staff explicit shift still wins on any day.
    custom = SimpleNamespace(shift_start="09:30")
    assert lf.start_for_day(date(2026, 10, 10), custom) == "09:30"


async def test_late_fine_uses_weekday_and_saturday_starts(
    db_session: AsyncSession, seed_data
):
    """A staff with no shift_start is judged against 08:00 on weekdays and 10:00
    on Saturday — reproduces the field case (the report previously read 0 late
    for everyone without an explicit shift)."""
    db_session.add_all([
        _slab(1, 5, 200, "#BDD7EE", "1 to 5"),
        _slab(6, 10, 500, "#FFEB84", "6 to 10"),
        _slab(11, None, 1000, "#FF6B6B", "11 or more"),
    ])
    dept = Department(branch_id=BRANCH, name="MSA-Teachers", id_range_start=1, id_range_end=50)
    db_session.add(dept)
    await db_session.flush()
    staff = Staff(
        branch_id=BRANCH, emp_code="91009", first_name="Dhiraj", last_name="Shrivastav",
        department_id=dept.id, shift_start=None,  # no explicit shift
    )
    db_session.add(staff)
    await db_session.flush()
    # IST = UTC+5:30. (date, utc_h, utc_m, expected local, weekday):
    #   Mon 05 Oct 08:04 IST (02:34 UTC) -> start 08:00 -> 4 late -> 200
    #   Tue 06 Oct 08:16 IST (02:46 UTC) -> start 08:00 -> 16 late -> 1000
    #   Sat 10 Oct 10:01 IST (04:31 UTC) -> start 10:00 -> 1 late -> 200
    #   Sat 17 Oct 10:30 IST (05:00 UTC) -> start 10:00 -> 30 late -> 1000
    for day, h, m in [(5, 2, 34), (6, 2, 46), (10, 4, 31), (17, 5, 0)]:
        db_session.add(StaffDailyAttendance(
            staff_id=staff.id, branch_id=BRANCH, attendance_date=date(2026, 10, day),
            first_in=datetime(2026, 10, day, h, m, tzinfo=UTC),
            day_status="PRESENT", signoff="COMPLETE", source="BIOMETRIC", work_minutes=400,
        ))
    await db_session.flush()

    _, data, _ = await lf.generate(
        db_session, branch_id=BRANCH, staff_id=staff.id,
        start=date(2026, 10, 5), end=date(2026, 10, 17), fmt="xlsx",
    )
    ws = load_workbook(io.BytesIO(data)).active
    rows = list(ws.iter_rows(values_only=True))
    hdr = next(i for i, r in enumerate(rows) if r[0] == "Date")
    data_rows = rows[hdr + 1: hdr + 5]
    assert [r[1] for r in data_rows] == ["08:00", "08:00", "10:00", "10:00"]  # Start
    assert [r[4] for r in data_rows] == [4, 16, 1, 30]                        # Late Minutes
    assert [r[5] for r in data_rows] == [200, 1000, 200, 1000]                # Fine
    total_row = rows[hdr + 5]
    assert total_row[3] == "Total Fine" and total_row[5] == 2400


async def test_late_fine_rejects_from_after_to(db_session: AsyncSession, seed_data):
    staff = await _seed(db_session)
    with pytest.raises(HTTPException) as ei:
        await lf.generate(db_session, branch_id=BRANCH, staff_id=staff.id,
                          start=date(2026, 10, 4), end=date(2026, 10, 1), fmt="pdf")
    assert ei.value.status_code == 422
