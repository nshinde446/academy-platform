"""Staff report engine — grid + typed list projections, Excel path (no headless
Chromium needed). Exercises the filter map and the report shapes."""

import uuid
from datetime import date, datetime, timezone

import pytest
from openpyxl import load_workbook
import io

from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.attendance.models.attendance_models import StaffDailyAttendance
from app.modules.attendance.services import staff_report_service as rpt
from app.modules.staff.models.staff_models import Department, Staff

BRANCH_A = uuid.UUID("00000000-0000-0000-0000-000000000001")


async def _seed(db: AsyncSession) -> Staff:
    dept = Department(branch_id=BRANCH_A, name="MSA-Teachers", id_range_start=1, id_range_end=50)
    db.add(dept)
    await db.flush()
    staff = Staff(
        branch_id=BRANCH_A, emp_code="1", first_name="Bhagvat", last_name="Dhesale",
        department_id=dept.id, shift_start="10:00", shift_end="20:00",
    )
    db.add(staff)
    await db.flush()

    def _utc(h, mi):
        return datetime(2026, 8, 3, h, mi, tzinfo=timezone.utc)

    # Day 1: present LATE (in 10:20 IST = 04:50 UTC, out 18:00 IST = 12:30 UTC).
    db.add(StaffDailyAttendance(
        staff_id=staff.id, branch_id=BRANCH_A, attendance_date=date(2026, 8, 3),
        first_in=_utc(4, 50), last_out=_utc(12, 30), day_status="LATE",
        signoff="COMPLETE", source="BIOMETRIC", work_minutes=460, ot_minutes=0,
    ))
    # Day 2: absent.
    db.add(StaffDailyAttendance(
        staff_id=staff.id, branch_id=BRANCH_A, attendance_date=date(2026, 8, 4),
        first_in=None, last_out=None, day_status="ABSENT",
        signoff="NA", source="SYSTEM", work_minutes=0, ot_minutes=0,
    ))
    # Day 3: mis-punch (in only).
    db.add(StaffDailyAttendance(
        staff_id=staff.id, branch_id=BRANCH_A, attendance_date=date(2026, 8, 5),
        first_in=_utc(4, 30), last_out=None, day_status="PRESENT",
        signoff="MISSING", source="BIOMETRIC", work_minutes=0, ot_minutes=0,
    ))
    await db.flush()
    return staff


async def _xlsx_first_col(data: bytes) -> list[str]:
    wb = load_workbook(io.BytesIO(data), read_only=True)
    ws = wb.active
    return [str(r[0].value) if r[0].value is not None else "" for r in ws.iter_rows()]


async def test_performance_grid_xlsx(db_session: AsyncSession, seed_data):
    await _seed(db_session)
    fn, data, mime = await rpt.generate(
        db_session, branch_id=BRANCH_A, kind="performance", frequency="monthly",
        department_ids=None, staff_ids=None,
        start=date(2026, 8, 3), end=date(2026, 8, 5), fmt="xlsx",
    )
    assert fn.endswith(".xlsx")
    assert mime == rpt.XLSX_MIME
    col_a = await _xlsx_first_col(data)
    # Staff header block present with the emp code + name.
    assert any("Bhagvat" in c and "1 ·" in c for c in col_a)


async def test_late_in_list_only_late_rows(db_session: AsyncSession, seed_data):
    await _seed(db_session)
    fn, data, _ = await rpt.generate(
        db_session, branch_id=BRANCH_A, kind="late_in", frequency="daily",
        department_ids=None, staff_ids=None,
        start=date(2026, 8, 3), end=date(2026, 8, 5), fmt="xlsx",
    )
    wb = load_workbook(io.BytesIO(data), read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    # Header block at top, then one data row (the LATE day).
    date_cells = [r[0] for r in rows if r[0] == "2026-08-03"]
    assert date_cells == ["2026-08-03"]
    assert not any(r[0] == "2026-08-04" for r in rows)  # absent day excluded


async def test_absent_list(db_session: AsyncSession, seed_data):
    await _seed(db_session)
    _, data, _ = await rpt.generate(
        db_session, branch_id=BRANCH_A, kind="absent", frequency="monthly",
        department_ids=None, staff_ids=None,
        start=date(2026, 8, 3), end=date(2026, 8, 5), fmt="xlsx",
    )
    wb = load_workbook(io.BytesIO(data), read_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    assert any(r[0] == "2026-08-04" for r in rows)
    assert not any(r[0] == "2026-08-03" for r in rows)


async def test_mis_punch_list(db_session: AsyncSession, seed_data):
    await _seed(db_session)
    _, data, _ = await rpt.generate(
        db_session, branch_id=BRANCH_A, kind="mis_punch", frequency="monthly",
        department_ids=None, staff_ids=None,
        start=date(2026, 8, 3), end=date(2026, 8, 5), fmt="xlsx",
    )
    rows = list(load_workbook(io.BytesIO(data), read_only=True).active.iter_rows(values_only=True))
    assert any(r[0] == "2026-08-05" for r in rows)
    assert not any(r[0] == "2026-08-03" for r in rows)


async def test_gps_report_is_not_applicable_note(db_session: AsyncSession, seed_data):
    await _seed(db_session)
    _, data, _ = await rpt.generate(
        db_session, branch_id=BRANCH_A, kind="gps_approved", frequency="daily",
        department_ids=None, staff_ids=None,
        start=date(2026, 8, 3), end=date(2026, 8, 3), fmt="xlsx",
    )
    col_a = await _xlsx_first_col(data)
    assert any("Not applicable" in c for c in col_a)


async def test_invalid_kind_rejected(db_session: AsyncSession, seed_data):
    await _seed(db_session)
    with pytest.raises(Exception):
        await rpt.generate(
            db_session, branch_id=BRANCH_A, kind="bogus", frequency="monthly",
            department_ids=None, staff_ids=None,
            start=date(2026, 8, 3), end=date(2026, 8, 5), fmt="xlsx",
        )
