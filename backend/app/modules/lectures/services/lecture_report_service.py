"""Lecture Report — a flat, per-lecture export (one row per scheduled lecture)
for a branch over a date range.

This is the client's "Teacher Productivity Report → Lecture Report" tab: every
lecture with its schedule-vs-actual, who delivered it, who recorded the entry,
and when. Times render **time-only** in the branch timezone; the date has its
own column. Duration can be shown as raw minutes or HH:MM.

Pure-ish: one gathering query + bulk name lookups, then Excel/PDF builders.
"""

from __future__ import annotations

import uuid
from datetime import date, datetime, timezone

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config.settings import get_settings
from app.modules.academic.models.academic_models import Subject, Topic
from app.modules.attendance.services import daily_service
from app.modules.attendance.services import staff_export_service as ex
from app.modules.attendance.time_utils import day_bounds, get_tz
from app.modules.auth.models.auth_models import Branch, User
from app.modules.batch.models.batch_models import Batch
from app.modules.classroom.models.classroom_models import Classroom
from app.modules.lectures.models.lecture_models import Lecture
from app.modules.teacher.models.teacher_models import Teacher

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MIME = "application/pdf"

HEADERS = [
    "Date", "Branch", "Batch", "Teacher", "Actual Teacher", "Subject", "Topic",
    "Classroom", "Delivery", "Scheduled Start", "Scheduled End",
    "Actual Start", "Actual End", "Duration", "Status",
    "Attendance Updated By", "Timestamp", "Late Start", "No-show Reason", "Notes",
]


def _aware(dt: datetime | None) -> datetime | None:
    if dt is not None and dt.tzinfo is None:
        return dt.replace(tzinfo=timezone.utc)
    return dt


def _t(dt: datetime | None, tz: str) -> str:
    """Time-only in the branch tz (the columns that used to show date+time)."""
    dt = _aware(dt)
    return dt.astimezone(get_tz(tz)).strftime("%I:%M %p").lstrip("0") if dt else ""


def _dt(dt: datetime | None, tz: str) -> str:
    """Full timestamp in the branch tz (the entry-made column)."""
    dt = _aware(dt)
    return dt.astimezone(get_tz(tz)).strftime("%d-%b-%Y %I:%M %p") if dt else ""


def _duration(l: Lecture, style: str) -> str:
    """Actual duration; falls back to scheduled length when actuals are absent.

    ``style='hhmm'`` -> "H:MM"; anything else -> raw minutes as a string.
    """
    mins = l.actual_duration_min
    if mins is None:
        a, b = _aware(l.actual_start), _aware(l.actual_end)
        if a and b:
            mins = max(0, int((b - a).total_seconds() // 60))
    if mins is None:
        a, b = _aware(l.scheduled_start), _aware(l.scheduled_end)
        mins = max(0, int((b - a).total_seconds() // 60)) if a and b else 0
    if style == "hhmm":
        return f"{mins // 60}:{mins % 60:02d}"
    return str(mins)


def _late(l: Lecture, tz: str) -> str:
    """Minutes late (actual start beyond scheduled start), blank if on time."""
    if not l.late_flag or not l.actual_start:
        return ""
    a, b = _aware(l.scheduled_start), _aware(l.actual_start)
    mins = max(0, int((b - a).total_seconds() // 60)) if a and b else 0
    return f"{mins} min" if mins else ""


def _title(v: str | None) -> str:
    return (v or "").replace("_", " ").title()


async def _name_maps(
    session: AsyncSession, lectures: list[Lecture]
) -> dict[str, dict[uuid.UUID, str]]:
    """Bulk-resolve the FKs on the lecture rows to display names."""
    teacher_ids = {l.teacher_id for l in lectures} | {
        l.actual_teacher_id for l in lectures if l.actual_teacher_id
    }
    batch_ids = {l.batch_id for l in lectures}
    subject_ids = {l.subject_id for l in lectures}
    topic_ids = {l.topic_id for l in lectures if l.topic_id}
    classroom_ids = {l.classroom_id for l in lectures if l.classroom_id}
    user_ids = {l.updated_by for l in lectures if l.updated_by}

    async def _pairs(model, ids, label):
        if not ids:
            return {}
        rows = (await session.execute(
            select(model.id, label).where(model.id.in_(ids))
        )).all()
        return {rid: name for rid, name in rows}

    teachers = {}
    if teacher_ids:
        for tid, fn, ln in (await session.execute(
            select(Teacher.id, Teacher.first_name, Teacher.last_name).where(
                Teacher.id.in_(teacher_ids)
            )
        )).all():
            teachers[tid] = f"{fn} {ln}".strip()

    users = {}
    if user_ids:
        for uid, fn, ln in (await session.execute(
            select(User.id, User.first_name, User.last_name).where(
                User.id.in_(user_ids)
            )
        )).all():
            users[uid] = f"{fn} {ln}".strip()

    return {
        "teacher": teachers,
        "user": users,
        "batch": await _pairs(Batch, batch_ids, Batch.name),
        "subject": await _pairs(Subject, subject_ids, Subject.name),
        "topic": await _pairs(Topic, topic_ids, Topic.name),
        "classroom": await _pairs(Classroom, classroom_ids, Classroom.name),
    }


async def _rows(
    session: AsyncSession, *, branch_id: uuid.UUID, start: date, end: date,
    tz: str, branch_name: str, duration_style: str,
    teacher_id: uuid.UUID | None, batch_id: uuid.UUID | None,
    subject_id: uuid.UUID | None,
) -> list[list[str]]:
    win_start, _ = day_bounds(start, tz)
    _, win_end = day_bounds(end, tz)
    q = select(Lecture).where(
        Lecture.branch_id == branch_id,
        Lecture.scheduled_start >= win_start,
        Lecture.scheduled_start < win_end,
        Lecture.is_deleted == False,  # noqa: E712
    )
    if teacher_id:
        q = q.where(
            (Lecture.actual_teacher_id == teacher_id)
            | ((Lecture.actual_teacher_id.is_(None)) & (Lecture.teacher_id == teacher_id))
        )
    if batch_id:
        q = q.where(Lecture.batch_id == batch_id)
    if subject_id:
        q = q.where(Lecture.subject_id == subject_id)
    lectures = list((await session.execute(
        q.order_by(Lecture.scheduled_start)
    )).scalars().all())

    m = await _name_maps(session, lectures)
    rows: list[list[str]] = []
    for l in lectures:
        d = _aware(l.scheduled_start).astimezone(get_tz(tz)).date()
        rows.append([
            d.strftime("%d-%b-%Y"),
            branch_name,
            m["batch"].get(l.batch_id, ""),
            m["teacher"].get(l.teacher_id, ""),
            m["teacher"].get(l.actual_teacher_id, "") if l.actual_teacher_id else "",
            m["subject"].get(l.subject_id, ""),
            m["topic"].get(l.topic_id, "") if l.topic_id else "",
            m["classroom"].get(l.classroom_id, "") if l.classroom_id else "",
            _title(l.delivery_mode),
            _t(l.scheduled_start, tz),
            _t(l.scheduled_end, tz),
            _t(l.actual_start, tz),
            _t(l.actual_end, tz),
            _duration(l, duration_style),
            _title(l.lecture_status),
            m["user"].get(l.updated_by, "") if l.updated_by else "",
            _dt(l.updated_at, tz),
            _late(l, tz),
            _title(l.no_show_reason) if l.no_show_reason else "",
            l.notes or "",
        ])
    return rows


async def generate(
    session: AsyncSession, *, branch_id: uuid.UUID, start: date, end: date,
    fmt: str, duration_style: str = "min",
    teacher_id: uuid.UUID | None = None, batch_id: uuid.UUID | None = None,
    subject_id: uuid.UUID | None = None,
) -> tuple[str, bytes, str]:
    tz = await daily_service.branch_timezone(session, branch_id)
    branch = (await session.execute(
        select(Branch).where(Branch.id == branch_id)
    )).scalar_one_or_none()
    branch_name = branch.name if branch else ""

    rows = await _rows(
        session, branch_id=branch_id, start=start, end=end, tz=tz,
        branch_name=branch_name, duration_style=duration_style,
        teacher_id=teacher_id, batch_id=batch_id, subject_id=subject_id,
    )
    brand = get_settings().ACADEMY_BRAND_NAME
    gen = ex.generated_stamp(tz)
    base = f"lecture-report-{start}-{end}"
    if fmt == "xlsx":
        return f"{base}.xlsx", _xlsx(brand, start, end, rows, gen), XLSX_MIME
    return f"{base}.pdf", await ex.render_html_to_pdf(
        _html(brand, start, end, rows, gen), landscape=True), PDF_MIME


# ── Builders ────────────────────────────────────────────────────────────────

def _html(brand: str, start: date, end: date, rows: list[list[str]], gen: str) -> str:
    ths = "".join(f"<th>{ex._esc(h)}</th>" for h in HEADERS)
    body_rows = "".join(
        "<tr>" + "".join(f"<td>{ex._esc(c)}</td>" for c in r) + "</tr>"
        for r in rows
    )
    empty = "" if rows else "<p class='sub'>No lectures in this range.</p>"
    body = (
        f"<h1>{ex._esc(brand)} — Lecture Report</h1>"
        f"<p class='sub'>{ex._period(start, end)} · {len(rows)} lectures</p>"
        f"<table><tr>{ths}</tr>{body_rows}</table>{empty}"
    )
    return ex._doc(body, gen)


def _xlsx(brand: str, start: date, end: date, rows: list[list[str]], gen: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Lecture Report"
    ws["A1"] = brand
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Lecture Report"
    ws["A3"] = ex._period(start, end)
    ws["A4"] = gen
    head = 5
    for c, h in enumerate(HEADERS, start=1):
        ws.cell(row=head, column=c, value=h).font = Font(bold=True)
    for i, r in enumerate(rows, start=1):
        for c, v in enumerate(r, start=1):
            ws.cell(row=head + i, column=c, value=v)
    widths = (12, 14, 16, 18, 18, 12, 16, 12, 10, 14, 14, 14, 14, 10, 12, 18, 20, 10, 16, 24)
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A6"
    return ex._xlsx_bytes(wb)
