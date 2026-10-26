"""PDF / Excel exports for the Monthly Payroll Attendance Report and the staff
ledger (§7/§11).

Reuses the shared renderers in ``staff_export_service`` (``_doc`` / ``_esc`` /
``render_html_to_pdf`` / ``_xlsx_bytes`` / ``generated_stamp``) so payroll
documents carry the same branding and the numbers match the on-screen report
(both are fed by :mod:`report_service`).
"""

from __future__ import annotations

import calendar
import uuid

from fastapi import HTTPException
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config.settings import get_settings
from app.modules.attendance.services import staff_export_service as ex
from app.modules.attendance.services.daily_service import branch_timezone
from app.modules.payroll.services import report_service

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MIME = "application/pdf"

_REPORT_HEADERS = [
    "Code", "Name", "Req.", "Present", "Half", "Paid L/H", "Net",
    "CF in", "CF out", "Uncovered", "Deferred", "Late Fine",
]
_LEDGER_HEADERS = ["Month", "Entry Type", "Days", "Running Balance", "By", "When"]


def _num(v: float) -> str:
    """Trim trailing .0 so 26.0 shows as 26 but 25.5 stays 25.5."""
    f = float(v)
    return str(int(f)) if f == int(f) else f"{f:g}"


def _month_label(year: int, month: int) -> str:
    return f"{calendar.month_name[month]} {year}" if 1 <= month <= 12 else f"{year}-{month}"


async def report(
    session: AsyncSession,
    *,
    branch_id: uuid.UUID,
    year: int,
    month: int,
    fmt: str,
    department_id: uuid.UUID | None = None,
    staff_id: uuid.UUID | None = None,
) -> tuple[str, bytes, str]:
    if fmt not in ("pdf", "xlsx"):
        raise HTTPException(status_code=422, detail="fmt must be pdf or xlsx")
    data = await report_service.monthly_report(
        session, branch_id=branch_id, year=year, month=month,
        department_id=department_id, staff_id=staff_id,
    )
    brand = get_settings().ACADEMY_BRAND_NAME
    period = _month_label(year, month)
    status = data["lifecycle"]
    base = f"Payroll_{year}-{month:02d}"
    gen = ex.generated_stamp(await branch_timezone(session, branch_id))

    if fmt == "xlsx":
        return f"{base}.xlsx", _report_xlsx(brand, period, status, data), XLSX_MIME
    return f"{base}.pdf", await ex.render_html_to_pdf(
        _report_html(brand, period, status, data, gen)
    ), PDF_MIME


def _report_row_cells(r: dict) -> list[str]:
    return [
        r["emp_code"], r["name"], _num(r["required_days"]), _num(r["present_days"]),
        str(r["half_days"]), _num(r["paid_leave_days"]), f'{r["net_days"]:+g}',
        _num(r["balance_in"]), _num(r["balance_out"]), _num(r["uncovered_days"]),
        _num(r["deferred_days"]), _num(r["late_fine_total"]),
    ]


def _report_html(brand: str, period: str, status: str, data: dict, gen: str) -> str:
    ths = "".join(f"<th>{ex._esc(h)}</th>" for h in _REPORT_HEADERS)
    body = "".join(
        "<tr>" + "".join(f"<td class='c'>{ex._esc(c)}</td>" for c in _report_row_cells(r)) + "</tr>"
        for r in data["rows"]
    )
    t = data["totals"]
    totals_row = (
        "<tr>"
        "<td colspan='3' style='text-align:right;font-weight:bold'>Totals</td>"
        f"<td class='c' style='font-weight:bold'>{_num(t['present_days'])}</td>"
        "<td colspan='5'></td>"
        f"<td class='c' style='font-weight:bold'>{_num(t['uncovered_days'])}</td>"
        "<td></td>"
        f"<td class='c' style='font-weight:bold'>{_num(t['late_fine_total'])}</td>"
        "</tr>"
        if data["rows"] else ""
    )
    empty = "" if data["rows"] else "<p class='sub'>No staff for this selection.</p>"
    doc = (
        f"<h1>{ex._esc(brand)} — Monthly Payroll Attendance Report</h1>"
        f"<p class='sub'>{ex._esc(period)} · Status: {ex._esc(status)} · "
        f"{t['staff_count']} staff</p>"
        f"<table><tr>{ths}</tr>{body}{totals_row}</table>{empty}"
    )
    return ex._doc(doc, gen)


def _report_xlsx(brand: str, period: str, status: str, data: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Payroll"
    ws["A1"] = brand
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Monthly Payroll Attendance Report"
    ws["A2"].font = Font(bold=True, size=12)
    ws["A3"] = f"{period} · Status: {status}"

    head = 5
    for c, h in enumerate(_REPORT_HEADERS, start=1):
        ws.cell(row=head, column=c, value=h).font = Font(bold=True)
    for i, r in enumerate(data["rows"], start=head + 1):
        for c, val in enumerate(_report_row_cells(r), start=1):
            ws.cell(row=i, column=c, value=val)
    if data["rows"]:
        t = data["totals"]
        tr = head + 1 + len(data["rows"])
        ws.cell(row=tr, column=2, value="Totals").font = Font(bold=True)
        ws.cell(row=tr, column=4, value=_num(t["present_days"])).font = Font(bold=True)
        ws.cell(row=tr, column=12, value=_num(t["late_fine_total"])).font = Font(bold=True)
    else:
        ws.cell(row=head + 1, column=1, value="No staff for this selection.")
    for c, w in enumerate((10, 22, 7, 9, 6, 9, 7, 7, 7, 10, 9, 10), start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    return ex._xlsx_bytes(wb)


# ── staff ledger export ──────────────────────────────────────────────────────

async def ledger(
    session: AsyncSession, *, branch_id: uuid.UUID, staff_id: uuid.UUID, fmt: str
) -> tuple[str, bytes, str]:
    if fmt not in ("pdf", "xlsx"):
        raise HTTPException(status_code=422, detail="fmt must be pdf or xlsx")
    data = await report_service.staff_ledger(session, branch_id=branch_id, staff_id=staff_id)
    brand = get_settings().ACADEMY_BRAND_NAME
    base = f"PayrollLedger_{data['emp_code']}"
    gen = ex.generated_stamp(await branch_timezone(session, branch_id))
    if fmt == "xlsx":
        return f"{base}.xlsx", _ledger_xlsx(brand, data), XLSX_MIME
    return f"{base}.pdf", await ex.render_html_to_pdf(_ledger_html(brand, data, gen)), PDF_MIME


def _ledger_row_cells(r: dict) -> list[str]:
    when = r["created_at"].strftime("%d-%m-%Y") if r.get("created_at") else "—"
    return [
        _month_label(r["year"], r["month"]), r["entry_type"], f'{r["days"]:+g}',
        _num(r["running_balance"]), r.get("created_by") or "—", when,
    ]


def _ledger_html(brand: str, data: dict, gen: str) -> str:
    ths = "".join(f"<th>{ex._esc(h)}</th>" for h in _LEDGER_HEADERS)
    body = "".join(
        "<tr>" + "".join(f"<td class='c'>{ex._esc(c)}</td>" for c in _ledger_row_cells(r)) + "</tr>"
        for r in data["rows"]
    )
    empty = "" if data["rows"] else "<p class='sub'>No ledger entries.</p>"
    doc = (
        f"<h1>{ex._esc(brand)} — Carry-Forward Ledger</h1>"
        f"<p class='sub'>Empcode: {ex._esc(data['emp_code'])} · {ex._esc(data['name'])} · "
        f"Current balance: {_num(data['current_balance'])}</p>"
        f"<table><tr>{ths}</tr>{body}</table>{empty}"
    )
    return ex._doc(doc, gen)


def _ledger_xlsx(brand: str, data: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Ledger"
    ws["A1"] = brand
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Carry-Forward Ledger"
    ws["A2"].font = Font(bold=True, size=12)
    ws["A3"] = f"Empcode: {data['emp_code']} · {data['name']} · Balance: {_num(data['current_balance'])}"
    head = 5
    for c, h in enumerate(_LEDGER_HEADERS, start=1):
        ws.cell(row=head, column=c, value=h).font = Font(bold=True)
    for i, r in enumerate(data["rows"], start=head + 1):
        for c, val in enumerate(_ledger_row_cells(r), start=1):
            ws.cell(row=i, column=c, value=val)
    for c, w in enumerate((16, 18, 8, 16, 20, 12), start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    return ex._xlsx_bytes(wb)
