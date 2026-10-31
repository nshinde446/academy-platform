"""Admin leave settings — per-year quota per leave type + WOF opening balances
(Teacher Portal Changes §3 + WOF). Super Admin only (enforced in the routes)."""

from __future__ import annotations

import uuid
from datetime import date, datetime, timezone

from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.staff.models.staff_models import Staff
from app.modules.teacher.models.teacher_models import (
    LeaveTypeQuota,
    Teacher,
    TeacherWofOpening,
)
from app.modules.teacher_portal.services import wof_service
from app.modules.teacher_portal.services.leave_service import QUOTA_TYPES


async def get_quotas(session: AsyncSession, branch_id: uuid.UUID, year: int) -> dict:
    rows = (await session.execute(
        select(LeaveTypeQuota).where(
            LeaveTypeQuota.branch_id == branch_id,
            LeaveTypeQuota.year == year,
            LeaveTypeQuota.is_deleted == False,  # noqa: E712
        )
    )).scalars().all()
    by_type = {r.leave_type: float(r.quota_days) for r in rows}
    return {"year": year, "quotas": {t: by_type.get(t) for t in QUOTA_TYPES}}


async def set_quotas(
    session: AsyncSession, *, branch_id: uuid.UUID, year: int,
    quotas: dict[str, float], user_id: uuid.UUID,
) -> dict:
    existing = {
        r.leave_type: r for r in (await session.execute(
            select(LeaveTypeQuota).where(
                LeaveTypeQuota.branch_id == branch_id,
                LeaveTypeQuota.year == year,
                LeaveTypeQuota.is_deleted == False,  # noqa: E712
            )
        )).scalars().all()
    }
    for lt, val in quotas.items():
        if lt not in QUOTA_TYPES or val is None:
            continue
        row = existing.get(lt)
        if row is None:
            session.add(LeaveTypeQuota(
                branch_id=branch_id, year=year, leave_type=lt,
                quota_days=val, created_by=user_id,
            ))
        else:
            row.quota_days = val
            row.updated_by = user_id
    await session.flush()
    return await get_quotas(session, branch_id, year)


async def _teachers_with_staff(
    session: AsyncSession, branch_id: uuid.UUID
) -> list[tuple[Teacher, Staff | None]]:
    teachers = (await session.execute(
        select(Teacher).where(
            Teacher.branch_id == branch_id, Teacher.is_deleted == False,  # noqa: E712
        )
    )).scalars().all()
    staff = {
        s.linked_teacher_id: s for s in (await session.execute(
            select(Staff).where(
                Staff.branch_id == branch_id, Staff.is_deleted == False,  # noqa: E712
                Staff.linked_teacher_id.isnot(None),
            )
        )).scalars().all()
    }
    return [(t, staff.get(t.id)) for t in teachers]


async def list_wof(session: AsyncSession, branch_id: uuid.UUID) -> list[dict]:
    openings = {
        o.teacher_id: o for o in (await session.execute(
            select(TeacherWofOpening).where(
                TeacherWofOpening.branch_id == branch_id,
                TeacherWofOpening.is_deleted == False,  # noqa: E712
            )
        )).scalars().all()
    }
    out: list[dict] = []
    for t, s in await _teachers_with_staff(session, branch_id):
        o = openings.get(t.id)
        bal = await wof_service.balance(session, s) if s is not None else None
        out.append({
            "teacher_id": str(t.id),
            "emp_code": s.emp_code if s else None,
            "name": f"{t.first_name} {t.last_name}".strip(),
            "opening_balance": float(o.opening_balance) if o else None,
            "as_of_date": o.as_of_date.isoformat() if o else None,
            "accrued": bal["accrued"] if bal else None,
            "current_balance": bal["balance"] if bal else None,
        })
    out.sort(key=lambda r: r["name"])
    return out


async def set_wof_opening(
    session: AsyncSession, *, branch_id: uuid.UUID, teacher_id: uuid.UUID,
    opening_balance: float, as_of_date: date, user_id: uuid.UUID,
) -> dict:
    t = (await session.execute(
        select(Teacher).where(Teacher.id == teacher_id, Teacher.is_deleted == False)  # noqa: E712
    )).scalar_one_or_none()
    if t is None or t.branch_id != branch_id:
        raise HTTPException(status_code=404, detail="Teacher not found in this branch")
    row = (await session.execute(
        select(TeacherWofOpening).where(
            TeacherWofOpening.teacher_id == teacher_id,
            TeacherWofOpening.is_deleted == False,  # noqa: E712
        )
    )).scalar_one_or_none()
    if row is None:
        session.add(TeacherWofOpening(
            teacher_id=teacher_id, branch_id=branch_id,
            opening_balance=opening_balance, as_of_date=as_of_date, created_by=user_id,
        ))
    else:
        row.opening_balance = opening_balance
        row.as_of_date = as_of_date
        row.updated_by = user_id
        row.updated_at = datetime.now(timezone.utc)
    await session.flush()
    return {"teacher_id": str(teacher_id), "opening_balance": opening_balance,
            "as_of_date": as_of_date.isoformat()}


# ── WOF roster report (PDF / Excel) ──────────────────────────────────────────

_WOF_HEADERS = ["Sr No", "Teacher", "Code", "Opening", "Earned", "WOF Balance", "As of"]


async def wof_report(
    session: AsyncSession, *, branch_id: uuid.UUID, fmt: str
) -> tuple[str, bytes, str]:
    """A roster of every teacher's Weekly-Off balance — opening + earned +
    current — as PDF or Excel (the system-generated equivalent of the client's
    'WOF Till …' sheet)."""
    from app.core.config.settings import get_settings
    from app.modules.attendance.services import staff_export_service as ex
    from app.modules.attendance.services.daily_service import branch_timezone

    if fmt not in ("pdf", "xlsx"):
        raise HTTPException(status_code=422, detail="fmt must be pdf or xlsx")

    rows = await list_wof(session, branch_id)
    # Only teachers with a seeded opening have a balance to report.
    rows = [r for r in rows if r["current_balance"] is not None]
    brand = get_settings().ACADEMY_BRAND_NAME
    gen = ex.generated_stamp(await branch_timezone(session, branch_id))

    def cells(i: int, r: dict) -> list[str]:
        opening = r["opening_balance"] or 0
        earned = r["accrued"] or 0
        return [
            str(i), r["name"], r["emp_code"] or "—",
            f"{opening:g}", f"{earned:g}", f"{r['current_balance']:g}",
            r["as_of_date"] or "—",
        ]

    base = "WOF_Balance_Report"
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "WOF Balance"
        ws["A1"] = brand
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = "Weekly-Off / Leave Balance — Teachers"
        ws["A2"].font = Font(bold=True, size=12)
        head = 4
        for c, h in enumerate(_WOF_HEADERS, start=1):
            ws.cell(row=head, column=c, value=h).font = Font(bold=True)
        for i, r in enumerate(rows, start=1):
            for c, val in enumerate(cells(i, r), start=1):
                ws.cell(row=head + i, column=c, value=val)
        for c, w in enumerate((6, 26, 10, 9, 9, 12, 12), start=1):
            ws.column_dimensions[get_column_letter(c)].width = w
        return f"{base}.xlsx", ex._xlsx_bytes(wb), \
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    ths = "".join(f"<th>{ex._esc(h)}</th>" for h in _WOF_HEADERS)
    body = "".join(
        "<tr>" + "".join(f"<td class='c'>{ex._esc(v)}</td>" for v in cells(i, r)) + "</tr>"
        for i, r in enumerate(rows, start=1)
    )
    empty = "" if rows else "<p class='sub'>No teacher WOF balances seeded yet.</p>"
    doc = (
        f"<h1>{ex._esc(brand)} — Weekly-Off / Leave Balance</h1>"
        f"<p class='sub'>Teachers · {len(rows)} shown</p>"
        f"<table><tr>{ths}</tr>{body}</table>{empty}"
    )
    return f"{base}.pdf", await ex.render_html_to_pdf(ex._doc(doc, gen)), "application/pdf"
