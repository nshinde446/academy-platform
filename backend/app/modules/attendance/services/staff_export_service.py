"""Staff attendance report builders — Excel (openpyxl) + PDF (headless Chromium).

Pure builders: take already-queried rows and return file bytes. Two shapes cover
the whole requirement grid:

* **grid** — the TeamOffice-style monthly/weekly performance block: one staff
  section with a day column per date carrying IN/OUT/WORK/OT and a Status line.
* **list** — a flat, filtered projection (Date · EmpCode · Name · Dept · In ·
  Out · Work · OT · Status) that serves every typed report (Present, Absent,
  Late-IN, Early-OUT, OT, Half-Day, Mis-Punch, In/Out, …); the caller pre-filters
  the rows and passes the title.

Reuses the student export scaffolding (PDF render, generation stamp, crest, CSS).
"""

from __future__ import annotations

import html
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.modules.attendance.services.attendance_export_service import (
    _period,
    _xlsx_bytes,
    generated_stamp,  # noqa: F401 — re-exported for the report service
    render_html_to_pdf,  # noqa: F401 — re-exported for the report service
)
from app.modules.attendance.services.attendance_export_service import (
    _logo_data_uri,
)

_BOLD = Font(bold=True)
_CENTER = Alignment(horizontal="center")
_HEAD_FILL = PatternFill("solid", fgColor="EEF1F5")
_STATUS_FILL = {
    "PRESENT": PatternFill("solid", fgColor="D7F0DB"),
    "LATE": PatternFill("solid", fgColor="FCEFC7"),
    "HALF_DAY": PatternFill("solid", fgColor="FCEFC7"),
    "ABSENT": PatternFill("solid", fgColor="F8D7DA"),
    "WO": PatternFill("solid", fgColor="E7E7EF"),
    "HOLIDAY": PatternFill("solid", fgColor="E7E7EF"),
}

# Short status codes for the grid (mirrors the e-TimeOffice P/A/WO glyphs).
STATUS_CODE = {
    "PRESENT": "P", "LATE": "L", "HALF_DAY": "½", "ABSENT": "A",
    "WO": "WO", "HOLIDAY": "H",
}

_LIST_HEADERS = [
    "Date", "Emp Code", "Name", "Department", "In", "Out", "Work", "OT",
    "Status", "Entry Type", "Marked By",
]

# How the stored source maps to the report's "Entry Type" column (PDF3).
_ENTRY_TYPE_LABEL = {
    "MANUAL": "Manual", "BIOMETRIC": "Biometric", "SYSTEM": "System",
}


def _entry_type_label(source: Any) -> str:
    s = "" if source is None else str(source)
    return _ENTRY_TYPE_LABEL.get(s, s)


def _esc(v: Any) -> str:
    return html.escape("" if v is None else str(v))


def fmt_minutes(minutes: int | None) -> str:
    """Minutes -> H:MM (0 -> 0:00), matching the report's Work/OT columns."""
    m = int(minutes or 0)
    return f"{m // 60}:{m % 60:02d}"


# ── PDF document shell (self-contained; crest embedded) ─────────────────────

_PDF_CSS = """
* { box-sizing: border-box; }
body { font: 12px -apple-system, Segoe UI, Roboto, sans-serif; color: #1a1a1a; margin: 0; }
h1 { font-size: 16px; margin: 0 0 2px; }
.sub { color: #555; font-size: 11px; margin: 0 0 10px; }
.genstamp { color: #888; margin-top: 12px; }
table { border-collapse: collapse; width: 100%; margin-bottom: 14px; }
th, td { border: 1px solid #cfd6df; padding: 3px 6px; font-size: 10.5px; }
th { background: #eef1f5; text-align: left; }
td.c, th.c { text-align: center; }
.staff-hd { background: #f5f7fa; font-weight: bold; padding: 4px 6px; border: 1px solid #cfd6df; }
.P { background: #d7f0db; } .L { background: #fcefc7; } .A { background: #f8d7da; }
.WO { background: #e7e7ef; } .H { background: #e7e7ef; }
"""


def _doc(body: str, generated: str = "") -> str:
    logo = _logo_data_uri()
    crest = f"<img src='{logo}' style='height:34px;float:right'>" if logo else ""
    foot = f"<p class='sub genstamp'>{_esc(generated)}</p>" if generated else ""
    return (
        f"<!DOCTYPE html><html><head><meta charset='utf-8'>"
        f"<style>{_PDF_CSS}</style></head><body>{crest}{body}{foot}</body></html>"
    )


# ── List report (every typed report) ───────────────────────────────────────

def list_html(
    *, brand: str, title: str, subtitle: str, rows: list[dict], generated: str = "",
) -> str:
    ths = "".join(f"<th>{_esc(h)}</th>" for h in _LIST_HEADERS)
    body_rows = "".join(
        "<tr>"
        f"<td>{_esc(r['date'])}</td>"
        f"<td class='c'>{_esc(r['emp_code'])}</td>"
        f"<td>{_esc(r['name'])}</td>"
        f"<td>{_esc(r['department'])}</td>"
        f"<td class='c'>{_esc(r['in'])}</td>"
        f"<td class='c'>{_esc(r['out'])}</td>"
        f"<td class='c'>{_esc(r['work'])}</td>"
        f"<td class='c'>{_esc(r['ot'])}</td>"
        f"<td class='c {_esc(STATUS_CODE.get(r['status'], ''))}'>{_esc(r['status'])}</td>"
        f"<td class='c'>{_esc(_entry_type_label(r.get('entry_type')))}</td>"
        f"<td>{_esc(r.get('marked_by') or '')}</td>"
        "</tr>"
        for r in rows
    )
    empty = "" if rows else "<p class='sub'>No matching records for this period.</p>"
    body = (
        f"<h1>{_esc(brand)} — {_esc(title)}</h1>"
        f"<p class='sub'>{_esc(subtitle)}</p>"
        f"<table><tr>{ths}</tr>{body_rows}</table>{empty}"
    )
    return _doc(body, generated)


def list_xlsx(
    *, brand: str, title: str, subtitle: str, rows: list[dict], generated: str = "",
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = brand
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = title
    ws["A2"].font = Font(bold=True, size=12)
    ws["A3"] = subtitle
    ws["A4"] = generated

    head = 5
    for c, h in enumerate(_LIST_HEADERS, start=1):
        cell = ws.cell(row=head, column=c, value=h)
        cell.font = _BOLD
        cell.fill = _HEAD_FILL
    for i, r in enumerate(rows, start=head + 1):
        ws.cell(row=i, column=1, value=r["date"])
        ws.cell(row=i, column=2, value=r["emp_code"]).alignment = _CENTER
        ws.cell(row=i, column=3, value=r["name"])
        ws.cell(row=i, column=4, value=r["department"])
        ws.cell(row=i, column=5, value=r["in"]).alignment = _CENTER
        ws.cell(row=i, column=6, value=r["out"]).alignment = _CENTER
        ws.cell(row=i, column=7, value=r["work"]).alignment = _CENTER
        ws.cell(row=i, column=8, value=r["ot"]).alignment = _CENTER
        cell = ws.cell(row=i, column=9, value=r["status"])
        cell.alignment = _CENTER
        if r["status"] in _STATUS_FILL:
            cell.fill = _STATUS_FILL[r["status"]]
        ws.cell(
            row=i, column=10, value=_entry_type_label(r.get("entry_type"))
        ).alignment = _CENTER
        ws.cell(row=i, column=11, value=r.get("marked_by") or "")
    for c, w in enumerate((12, 10, 22, 20, 8, 8, 8, 8, 12, 12, 24), start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    return _xlsx_bytes(wb)


# ── Performance grid (TeamOffice-style monthly/weekly block) ────────────────

def grid_html(
    *, brand: str, title: str, subtitle: str, dates: list[date],
    staff_blocks: list[dict], generated: str = "",
) -> str:
    """One block per staff member: totals line + a table with a column per date
    carrying IN / OUT / WORK / OT / Status."""
    day_ths = "".join(f"<th class='c'>{d.day}/{d.month}</th>" for d in dates)
    blocks = []
    for b in staff_blocks:
        cells = b["cells"]  # aligned to dates
        def _row(field: str) -> str:
            return "".join(f"<td class='c'>{_esc(c.get(field, ''))}</td>" for c in cells)
        status_row = "".join(
            f"<td class='c {_esc(STATUS_CODE.get(c.get('status'), ''))}'>"
            f"{_esc(STATUS_CODE.get(c.get('status'), ''))}</td>"
            for c in cells
        )
        blocks.append(
            f"<div class='staff-hd'>{_esc(b['emp_code'])} · {_esc(b['name'])} · "
            f"{_esc(b['department'])}"
            f" — Present {b['present']} · Absent {b['absent']} · WO {b['wo']} · "
            f"Work+OT {b['work_ot']}</div>"
            f"<table><tr><th>Field</th>{day_ths}</tr>"
            f"<tr><th>In</th>{_row('in')}</tr>"
            f"<tr><th>Out</th>{_row('out')}</tr>"
            f"<tr><th>Work</th>{_row('work')}</tr>"
            f"<tr><th>OT</th>{_row('ot')}</tr>"
            f"<tr><th>Status</th>{status_row}</tr>"
            f"</table>"
        )
    empty = "" if staff_blocks else "<p class='sub'>No staff match the selection.</p>"
    body = (
        f"<h1>{_esc(brand)} — {_esc(title)}</h1>"
        f"<p class='sub'>{_esc(subtitle)}</p>{''.join(blocks)}{empty}"
    )
    return _doc(body, generated)


def grid_xlsx(
    *, brand: str, title: str, subtitle: str, dates: list[date],
    staff_blocks: list[dict], generated: str = "",
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Performance"
    ws["A1"] = brand
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = title
    ws["A2"].font = Font(bold=True, size=12)
    ws["A3"] = subtitle
    ws["A4"] = generated

    row = 6
    for b in staff_blocks:
        ws.cell(
            row=row, column=1,
            value=(
                f"{b['emp_code']} · {b['name']} · {b['department']} — "
                f"Present {b['present']} · Absent {b['absent']} · WO {b['wo']} · "
                f"Work+OT {b['work_ot']}"
            ),
        ).font = _BOLD
        row += 1
        ws.cell(row=row, column=1, value="Field").font = _BOLD
        for j, d in enumerate(dates):
            c = ws.cell(row=row, column=2 + j, value=f"{d.day}/{d.month}")
            c.font = _BOLD
            c.fill = _HEAD_FILL
            c.alignment = _CENTER
        for field in ("in", "out", "work", "ot", "status"):
            row += 1
            ws.cell(row=row, column=1, value=field.capitalize()).font = _BOLD
            for j, cell in enumerate(b["cells"]):
                val = STATUS_CODE.get(cell.get("status"), "") if field == "status" else cell.get(field, "")
                out = ws.cell(row=row, column=2 + j, value=val)
                out.alignment = _CENTER
                if field == "status" and cell.get("status") in _STATUS_FILL:
                    out.fill = _STATUS_FILL[cell["status"]]
        row += 2  # blank spacer between staff blocks

    ws.column_dimensions["A"].width = 10
    for j in range(len(dates)):
        ws.column_dimensions[get_column_letter(2 + j)].width = 6
    return _xlsx_bytes(wb)


# ── Date-Wise Attendance (Staff Monthly Attendance Report) ──────────────────
#
# One row per staff member: Empcode · Name · a cell per calendar day carrying a
# single status code (P/A/L/HD/WO/HL) · six totals (P/A/HD/L/HL/WO). L/HD/HL
# cells are highlighted red/blue/grey in both PDF and Excel. A legend and the
# client's note print below the table. Day columns stay legible in landscape.

DATEWISE_TOTAL_CODES = ["P", "A", "HD", "L", "HL", "WO"]

DATEWISE_LEGEND = [
    ("P", "Present"), ("A", "Absent"), ("L", "Late"),
    ("HD", "Half Day"), ("WO", "Weekly Off"), ("HL", "Holiday"),
]

# The single client-facing note that prints on the report (open-item default:
# note 1 prints; notes 2 & 3 are implemented as functional requirements).
DATEWISE_NOTE = "1. This is for your reference."

# Cell highlights for the coded day cells (red / blue / grey light shades).
_DATEWISE_HTML_BG = {"L": "#F8C9C9", "HD": "#C9DAF8", "HL": "#D9D9D9"}
_DATEWISE_XLSX_FILL = {
    "L": PatternFill("solid", fgColor="F8C9C9"),
    "HD": PatternFill("solid", fgColor="C9DAF8"),
    "HL": PatternFill("solid", fgColor="D9D9D9"),
}


def _datewise_day_heads(dates: list[date]) -> list[str]:
    """Day-column headers: the day number, or ``d/m`` when the range spans
    more than one month so the day stays unambiguous."""
    spans_months = len({(d.year, d.month) for d in dates}) > 1
    return [f"{d.day}/{d.month}" if spans_months else str(d.day) for d in dates]


def datewise_html(
    *, brand: str, title: str, header_label: str, dates: list[date],
    rows: list[dict], generated: str = "",
) -> str:
    """``rows`` items: {emp_code, name, cells: [code,…] aligned to ``dates``,
    totals: {code: count}}. ``header_label`` is the "Month - …" / "Period - …"
    banner; ``title`` is the fixed "Date Wise Attendance"."""
    heads = _datewise_day_heads(dates)
    day_ths = "".join(f"<th class='c'>{_esc(h)}</th>" for h in heads)
    total_ths = "".join(f"<th class='c'>{_esc(code)}</th>" for code in DATEWISE_TOTAL_CODES)
    body_rows = []
    for r in rows:
        day_tds = "".join(
            (
                f"<td class='c' style='background:{_DATEWISE_HTML_BG[code]}'>{_esc(code)}</td>"
                if code in _DATEWISE_HTML_BG
                else f"<td class='c'>{_esc(code)}</td>"
            )
            for code in r["cells"]
        )
        total_tds = "".join(
            f"<td class='c'>{r['totals'].get(code, 0)}</td>" for code in DATEWISE_TOTAL_CODES
        )
        body_rows.append(
            f"<tr><td class='c'>{_esc(r['emp_code'])}</td>"
            f"<td>{_esc(r['name'])}</td>{day_tds}{total_tds}</tr>"
        )
    empty = "" if rows else "<p class='sub'>No staff match the selection.</p>"
    legend = " &nbsp;·&nbsp; ".join(f"<b>{_esc(c)}</b> = {_esc(m)}" for c, m in DATEWISE_LEGEND)
    body = (
        f"<h1>{_esc(brand)} — {_esc(title)}</h1>"
        f"<p class='sub'>{_esc(header_label)}</p>"
        f"<table style='font-size:8.5px'><tr>"
        f"<th>Empcode</th><th>Name</th>{day_ths}{total_ths}</tr>"
        f"{''.join(body_rows)}</table>{empty}"
        f"<div class='sub' style='margin-top:10px'><b>Legend:</b> {legend}</div>"
        f"<div class='sub' style='margin-top:6px'><b>Notes:</b><div>{_esc(DATEWISE_NOTE)}</div></div>"
    )
    return _doc(body, generated)


def datewise_xlsx(
    *, brand: str, title: str, header_label: str, dates: list[date],
    rows: list[dict], generated: str = "",
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Date Wise Attendance"
    ws["A1"] = brand
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = title
    ws["A2"].font = Font(bold=True, size=12)
    ws["A3"] = header_label
    ws["A4"] = generated

    heads = _datewise_day_heads(dates)
    head_row = 6
    headers = ["Empcode", "Name", *heads, *DATEWISE_TOTAL_CODES]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=head_row, column=c, value=h)
        cell.font = _BOLD
        cell.fill = _HEAD_FILL
        cell.alignment = _CENTER
    first_day_col = 3  # columns 1,2 = Empcode, Name
    for i, r in enumerate(rows, start=head_row + 1):
        ws.cell(row=i, column=1, value=r["emp_code"]).alignment = _CENTER
        ws.cell(row=i, column=2, value=r["name"])
        for j, code in enumerate(r["cells"]):
            cell = ws.cell(row=i, column=first_day_col + j, value=code)
            cell.alignment = _CENTER
            if code in _DATEWISE_XLSX_FILL:
                cell.fill = _DATEWISE_XLSX_FILL[code]
        total_col = first_day_col + len(dates)
        for k, code in enumerate(DATEWISE_TOTAL_CODES):
            ws.cell(
                row=i, column=total_col + k, value=r["totals"].get(code, 0)
            ).alignment = _CENTER

    # Legend + note below the table.
    gap = head_row + len(rows) + 2
    ws.cell(row=gap, column=1, value="Legend").font = _BOLD
    for n, (code, meaning) in enumerate(DATEWISE_LEGEND, start=1):
        ws.cell(row=gap + n, column=1, value=code).font = _BOLD
        ws.cell(row=gap + n, column=2, value=meaning)
    note_row = gap + len(DATEWISE_LEGEND) + 2
    ws.cell(row=note_row, column=1, value="Notes").font = _BOLD
    ws.cell(row=note_row + 1, column=1, value=DATEWISE_NOTE)

    ws.freeze_panes = ws.cell(row=head_row + 1, column=first_day_col)
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 22
    for j in range(len(dates) + len(DATEWISE_TOTAL_CODES)):
        ws.column_dimensions[get_column_letter(first_day_col + j)].width = 4.5
    return _xlsx_bytes(wb)
