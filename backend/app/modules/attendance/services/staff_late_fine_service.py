"""Staff Late Fine Report (requirement doc).

Late minutes = IN Time − First Lecture Start Time for the same date (minute
precision, seconds ignored; never negative). The fine comes from configurable
slabs (``late_fine_slabs``). One shared function (:func:`late_minutes` +
:func:`fine_for`) feeds both the data and every export so the numbers always
match. No forgiveness rule until the client confirms (every late day is fined
strictly per the slabs).
"""

from __future__ import annotations

import re
import uuid
from datetime import date, datetime, timezone

from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config.settings import get_settings
from app.modules.attendance.models.attendance_models import (
    LateFineSlab,
    StaffDailyAttendance,
)
from app.modules.attendance.services import daily_service
from app.modules.attendance.services import staff_export_service as ex
from app.modules.attendance.time_utils import get_tz
from app.modules.staff.models.staff_models import Staff

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MIME = "application/pdf"


def _parse_hhmm(value: str | None) -> tuple[int, int] | None:
    """Parse a wall-clock time the staff's First Lecture Start Time is entered
    as. Tolerant of the separator the academy actually uses: ``08:00``, ``08.00``
    (dot), ``8-00``, ``0800`` and ``800`` all parse to (8, 0). Returns None only
    when no valid hour/minute can be read."""
    if value is None:
        return None
    digits = re.findall(r"\d+", str(value))
    if len(digits) >= 2:
        h, m = int(digits[0]), int(digits[1])
    elif len(digits) == 1 and len(digits[0]) in (3, 4):
        d = digits[0].zfill(4)  # "800" -> "0800"
        h, m = int(d[:2]), int(d[2:])
    else:
        return None
    return (h, m) if 0 <= h < 24 and 0 <= m < 60 else None


def late_minutes(start_hhmm: str | None, in_local: datetime | None) -> int:
    """Whole minutes late = local IN minute-of-day − shift-start minute-of-day,
    floored at 0 (early/on-time = 0). Seconds are ignored."""
    hm = _parse_hhmm(start_hhmm)
    if hm is None or in_local is None:
        return 0
    start = hm[0] * 60 + hm[1]
    arrived = in_local.hour * 60 + in_local.minute
    return max(0, arrived - start)


def fine_for(minutes: int, slabs: list[LateFineSlab]) -> tuple[float, LateFineSlab | None]:
    """The fine + matching slab for ``minutes`` late. 0 minutes → no fine."""
    if minutes <= 0:
        return 0.0, None
    for slab in slabs:
        if minutes >= slab.min_minutes and (slab.max_minutes is None or minutes <= slab.max_minutes):
            return slab.fine_amount, slab
    return 0.0, None


def start_for_day(day: date, staff: Staff) -> str:
    """The "First Lecture Start Time" to judge a given day against. A staff
    member's own ``shift_start`` wins when set; otherwise the institute shift —
    08:00 on weekdays (Mon–Fri), 10:00 on the weekend (Sat/Sun)."""
    if staff.shift_start:
        return staff.shift_start
    s = get_settings()
    return s.LATE_FINE_WEEKEND_START if day.weekday() >= 5 else s.LATE_FINE_WEEKDAY_START


def _hhmm(minutes: int) -> str:
    return f"{minutes // 60:02d}:{minutes % 60:02d}"


def _fmt_in(in_local: datetime | None) -> str:
    return in_local.strftime("%H:%M") if in_local else "—"


async def _load_slabs(session: AsyncSession) -> list[LateFineSlab]:
    return list((await session.execute(
        select(LateFineSlab).where(LateFineSlab.is_deleted == False)  # noqa: E712
        .order_by(LateFineSlab.min_minutes)
    )).scalars().all())


async def generate(
    session: AsyncSession, *, branch_id: uuid.UUID, staff_id: uuid.UUID,
    start: date, end: date, fmt: str,
) -> tuple[str, bytes, str]:
    """The Staff Late Fine Report for one staff member over [start, end], as PDF
    or Excel. Default rows = every date in the range with an IN punch."""
    if fmt not in ("pdf", "xlsx"):
        raise HTTPException(status_code=422, detail="fmt must be pdf or xlsx")
    if start > end:
        raise HTTPException(status_code=422, detail="From date must not be after To date")

    staff = (await session.execute(
        select(Staff).where(Staff.id == staff_id, Staff.is_deleted == False)  # noqa: E712
    )).scalar_one_or_none()
    if staff is None or staff.branch_id != branch_id:
        raise HTTPException(status_code=404, detail="Staff not found in this branch")

    tz_name = await daily_service.branch_timezone(session, branch_id)
    tz = get_tz(tz_name)
    slabs = await _load_slabs(session)

    days = (await session.execute(
        select(StaffDailyAttendance).where(
            StaffDailyAttendance.staff_id == staff_id,
            StaffDailyAttendance.branch_id == branch_id,
            StaffDailyAttendance.attendance_date >= start,
            StaffDailyAttendance.attendance_date <= end,
            StaffDailyAttendance.is_deleted == False,  # noqa: E712
            StaffDailyAttendance.first_in.isnot(None),
        ).order_by(StaffDailyAttendance.attendance_date)
    )).scalars().all()

    rows: list[dict] = []
    for d in days:
        in_utc = d.first_in
        if in_utc is not None and in_utc.tzinfo is None:
            in_utc = in_utc.replace(tzinfo=timezone.utc)  # stored UTC; SQLite drops tzinfo
        in_local = in_utc.astimezone(tz) if in_utc else None
        day_start = start_for_day(d.attendance_date, staff)
        mins = late_minutes(day_start, in_local)
        fine, slab = fine_for(mins, slabs)
        rows.append({
            "date": d.attendance_date.strftime("%d-%m-%Y"),
            "start": day_start,
            "in": _fmt_in(in_local),
            "late_in": _hhmm(mins),
            "late_minutes": mins,
            "fine": fine,
            "color": slab.color_hex if slab else None,
        })

    total = round(sum(r["fine"] for r in rows), 2)
    name = f"{staff.first_name} {staff.last_name}".strip()
    period = f"{start.strftime('%d-%m-%Y')} to {end.strftime('%d-%m-%Y')}"
    base = f"LateFine_{staff.emp_code}_{start.isoformat()}_to_{end.isoformat()}"
    brand = get_settings().ACADEMY_BRAND_NAME
    gen = ex.generated_stamp(tz_name)
    report = {
        "brand": brand, "emp_code": staff.emp_code, "name": name,
        "period": period, "rows": rows, "total": total,
    }

    if fmt == "xlsx":
        return f"{base}.xlsx", _xlsx(report), XLSX_MIME
    return f"{base}.pdf", await ex.render_html_to_pdf(_html(report, gen)), PDF_MIME


# ── Renderers ────────────────────────────────────────────────────────────────

_HEADERS = ["Date", "First Lecture Start Time", "IN Time", "Late IN",
            "Late Minutes", "Fine Amount"]


def _html(r: dict, gen: str) -> str:
    ths = "".join(f"<th>{ex._esc(h)}</th>" for h in _HEADERS)
    body = "".join(
        "<tr>"
        f"<td class='c'>{ex._esc(row['date'])}</td>"
        f"<td class='c'>{ex._esc(row['start'])}</td>"
        f"<td class='c'>{ex._esc(row['in'])}</td>"
        f"<td class='c'>{ex._esc(row['late_in'])}</td>"
        f"<td class='c'"
        + (f" style='background:{row['color']}'" if row["color"] else "")
        + f">{row['late_minutes']}</td>"
        f"<td class='c'>{int(row['fine'])}</td>"
        "</tr>"
        for row in r["rows"]
    )
    empty = "" if r["rows"] else "<p class='sub'>No attendance data for this selection.</p>"
    total_row = (
        f"<tr><td colspan='5' style='text-align:right;font-weight:bold'>Total Fine</td>"
        f"<td class='c' style='font-weight:bold'>{int(r['total'])}</td></tr>"
        if r["rows"] else ""
    )
    doc = (
        f"<h1>{ex._esc(r['brand'])} — Staff Late Fine Report</h1>"
        f"<p class='sub'>Empcode: {ex._esc(r['emp_code'])} · {ex._esc(r['name'])} · {ex._esc(r['period'])}</p>"
        f"<table><tr>{ths}</tr>{body}{total_row}</table>{empty}"
    )
    return ex._doc(doc, gen)


def _xlsx(r: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Late Fine"
    ws["A1"] = r["brand"]
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Staff Late Fine Report"
    ws["A2"].font = Font(bold=True, size=12)
    ws["A3"] = f"Empcode: {r['emp_code']} · {r['name']}"
    ws["A4"] = r["period"]

    head = 6
    for c, h in enumerate(_HEADERS, start=1):
        ws.cell(row=head, column=c, value=h).font = Font(bold=True)
    for i, row in enumerate(r["rows"], start=head + 1):
        ws.cell(row=i, column=1, value=row["date"])
        ws.cell(row=i, column=2, value=row["start"])
        ws.cell(row=i, column=3, value=row["in"])
        ws.cell(row=i, column=4, value=row["late_in"])
        cell = ws.cell(row=i, column=5, value=row["late_minutes"])
        if row["color"]:
            hexval = row["color"].lstrip("#").upper()
            cell.fill = PatternFill("solid", fgColor=hexval)
        ws.cell(row=i, column=6, value=int(row["fine"]))
    if r["rows"]:
        tr = head + 1 + len(r["rows"])
        ws.cell(row=tr, column=4, value="Total Fine").font = Font(bold=True)
        ws.cell(row=tr, column=6, value=int(r["total"])).font = Font(bold=True)
    else:
        ws.cell(row=head + 1, column=1, value="No attendance data for this selection.")

    for c, w in enumerate((12, 22, 10, 10, 14, 14), start=1):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(c)].width = w
    return ex._xlsx_bytes(wb)
