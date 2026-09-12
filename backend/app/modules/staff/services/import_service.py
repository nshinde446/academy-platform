import csv
import io
import uuid
from typing import Any

from fastapi import HTTPException, UploadFile, status
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.audit.services import audit_service
from app.modules.staff.models.staff_models import Department
from app.modules.staff.repositories import staff_repository
from app.modules.staff.services import staff_service

# CSV/Excel column header -> Staff field. Mirrors the reference employee sheet
# (EmpCode, Name, Department, Designation).
COLUMN_MAPPING = {
    "empcode": "emp_code",
    "emp code": "emp_code",
    "employee code": "emp_code",
    "employee id": "emp_code",
    "code": "emp_code",
    "name": "name",
    "first name": "first_name",
    "first_name": "first_name",
    "last name": "last_name",
    "last_name": "last_name",
    "department": "department",
    "dept": "department",
    "designation": "designation",
    "role": "designation",
    "email": "email",
    "phone": "phone",
}

# Optional-title tokens stripped off the front of a name ("Mr. Ram" -> Ram).
_TITLES = {"mr", "mr.", "mrs", "mrs.", "ms", "ms.", "dr", "dr."}


def _normalize(header: str) -> str:
    return header.strip().lower()


def _split_name(value: str) -> tuple[str | None, str, str]:
    """-> (title, first_name, last_name)."""
    parts = value.strip().split()
    title: str | None = None
    if parts and parts[0].lower() in _TITLES:
        title = parts[0].rstrip(".")
        parts = parts[1:]
    if not parts:
        return title, "", ""
    if len(parts) == 1:
        return title, parts[0], ""
    return title, parts[0], " ".join(parts[1:])


def _parse_csv(content: bytes) -> list[dict[str, Any]]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return [dict(row) for row in reader]


def _parse_xlsx(content: bytes) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = wb.active
    if ws is None:
        return []
    rows = ws.iter_rows(values_only=True)
    try:
        headers = [str(h) if h is not None else "" for h in next(rows)]
    except StopIteration:
        return []
    out: list[dict[str, Any]] = []
    for row in rows:
        record = {
            headers[i]: ("" if row[i] is None else str(row[i]))
            for i in range(min(len(headers), len(row)))
        }
        if any(v for v in record.values()):
            out.append(record)
    return out


def _map_row(row: dict[str, Any]) -> dict[str, Any]:
    mapped: dict[str, Any] = {}
    for raw_key, raw_value in row.items():
        if raw_key is None:
            continue
        field = COLUMN_MAPPING.get(_normalize(str(raw_key)))
        if field is None:
            continue
        value = (raw_value or "").strip() if isinstance(raw_value, str) else raw_value
        if value in ("", None):
            continue
        mapped[field] = str(value).strip()

    name_value = mapped.pop("name", None)
    if name_value:
        title, first, last = _split_name(str(name_value))
        mapped.setdefault("title", title)
        mapped.setdefault("first_name", first)
        mapped.setdefault("last_name", last)
    return mapped


async def import_staff(
    session: AsyncSession,
    file: UploadFile,
    branch_id: uuid.UUID,
    current_user_id: uuid.UUID,
    ip_address: str | None = None,
) -> dict[str, Any]:
    content = await file.read()
    filename = (file.filename or "").lower()
    if filename.endswith(".csv"):
        rows = _parse_csv(content)
    elif filename.endswith((".xlsx", ".xls")):
        rows = _parse_xlsx(content)
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file type. Use .csv or .xlsx",
        )

    await staff_service.ensure_departments(session, branch_id)
    # Department name -> row (case-insensitive).
    dept_by_name = {
        d.name.strip().lower(): d
        for d in await staff_repository.list_departments(session, branch_id)
    }

    imported = 0
    skipped = 0
    errors: list[str] = []

    for idx, row in enumerate(rows, start=2):  # row 1 is header
        mapped = _map_row(row)
        if not mapped.get("first_name"):
            skipped += 1
            errors.append(f"Row {idx}: missing 'Name'")
            continue
        dept_name = str(mapped.pop("department", "") or "").strip().lower()
        dept: Department | None = dept_by_name.get(dept_name)
        if dept is None:
            skipped += 1
            errors.append(f"Row {idx}: unknown department '{dept_name or '(blank)'}'")
            continue

        emp_code = mapped.pop("emp_code", None)
        try:
            if emp_code:
                emp_code = str(emp_code).strip()
                if await staff_repository.get_by_emp_code(session, branch_id, emp_code):
                    skipped += 1
                    errors.append(f"Row {idx}: emp_code {emp_code} already exists")
                    continue
                is_legacy = not staff_service._code_in_range(emp_code, dept)
            else:
                emp_code = await staff_service.allocate_emp_code(session, branch_id, dept)
                is_legacy = False

            await staff_repository.create(
                session,
                branch_id=branch_id,
                department_id=dept.id,
                emp_code=emp_code,
                is_legacy_code=is_legacy,
                title=mapped.get("title"),
                first_name=mapped["first_name"],
                last_name=mapped.get("last_name", ""),
                designation=mapped.get("designation"),
                email=mapped.get("email"),
                phone=mapped.get("phone"),
            )
            imported += 1
        except Exception as exc:  # noqa: BLE001
            skipped += 1
            errors.append(f"Row {idx}: {exc}")

    await audit_service.log_action(
        session,
        user_id=current_user_id,
        action="IMPORT",
        table_name="staff",
        record_id=uuid.uuid4(),
        new_values={"imported": imported, "skipped": skipped},
        ip_address=ip_address,
        branch_id=branch_id,
    )
    return {"imported": imported, "skipped": skipped, "errors": errors}
